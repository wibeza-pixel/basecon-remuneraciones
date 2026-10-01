"""Documentos Word: contrato, liquidación, comprobante de feriado y finiquito."""
from __future__ import annotations

import json
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from . import config as C
from .calculos import numero_a_palabras

def _fmt_clp(n):
    """Formato pesos chilenos: 1.234.567"""
    try:
        return f"{int(round(float(n or 0))):,}".replace(",", ".")
    except Exception:
        return "0"


def _set_run(paragraph, text, bold=False, size=9, center=False):
    paragraph.clear() if hasattr(paragraph, "clear") else None
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(str(text))
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.bold = bold
    return run


def _cell(cell, text, bold=False, size=9, align="left"):
    """Escribe en celda limpiando párrafos previos."""
    cell.text = ""
    p = cell.paragraphs[0]
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(str(text if text is not None else ""))
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.bold = bold


def _set_table_borders(table, color="000000", sz="4"):
    """Bordes finos tipo formulario (no grilla gruesa de Excel)."""
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    # quitar bordes previos
    for child in list(tblPr):
        if child.tag == qn("w:tblBorders"):
            tblPr.remove(child)
    tblPr.append(borders)



def _cell_borders(cell, **lados):
    """Bordes por celda: lados = top/bottom/left/right con valores 'single' o None (sin borde)."""
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag == qn("w:tcBorders"):
            tcPr.remove(child)
    b = OxmlElement("w:tcBorders")
    for lado in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{lado}")
        val = lados.get(lado)
        el.set(qn("w:val"), "single" if val else "nil")
        if val:
            el.set(qn("w:sz"), str(val if isinstance(val, int) else 6))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "000000")
        b.append(el)
    tcPr.append(b)


def _sin_bordes(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    for child in list(tblPr):
        if child.tag == qn("w:tblBorders"):
            tblPr.remove(child)
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        b.append(el)
    tblPr.append(b)


def _anchos(table, anchos_cm):
    table.autofit = False
    for i, w in enumerate(anchos_cm):
        table.columns[i].width = Cm(w)
        for r in table.rows:
            r.cells[i].width = Cm(w)


def _txt(cell, texto, bold=False, size=9, align="left", espaciado=0, font="Calibri"):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT,
                   "center": WD_ALIGN_PARAGRAPH.CENTER}[align]
    r = p.add_run("" if texto is None else str(texto))
    r.bold = bold
    r.font.name = font
    r.font.size = Pt(size)
    if espaciado:
        rPr = r._r.get_or_add_rPr()
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:val"), str(espaciado))
        rPr.append(sp)
    return r


def generar_liquidacion_docx(empresa, trabajador, liq, periodo, ruta, indicadores=None):
    """
    Liquidación de remuneraciones en el formato de la oficina (modelo NUT EXPORT, agosto 2026):
    encabezado de la empresa a la izquierda, título espaciado y subrayado, ficha del trabajador en recuadro,
    columnas HABERES | DESCUENTOS con subtotales, totales, líquido, monto en palabras y firma.
    """
    FONT = "Calibri"
    doc = Document()
    for sec in doc.sections:
        sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
        sec.top_margin = Cm(1.3)
        sec.bottom_margin = Cm(1.2)
        sec.left_margin = Cm(1.5)
        sec.right_margin = Cm(1.5)
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(9)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    ind = indicadores or {}
    uf = float(ind.get("uf") or 0)
    tope_afp = float(ind.get("tope_afp") or 0)
    tope_afc = float(ind.get("tope_afc") or 0)

    def par(texto, bold=False, size=9, align="left", subrayado=False, espaciado=0, after=2):
        p = doc.add_paragraph()
        p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                       "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(after)
        r = p.add_run(texto)
        r.bold, r.underline = bold, subrayado
        r.font.name, r.font.size = FONT, Pt(size)
        if espaciado:
            sp = OxmlElement("w:spacing")
            sp.set(qn("w:val"), str(espaciado))
            r._r.get_or_add_rPr().append(sp)
        return p

    # ----- Encabezado de la empresa -----
    par(str(empresa.get("razon_social") or "").upper(), bold=True, size=11, after=3)
    par(str(empresa.get("rut") or ""), bold=True, size=8)
    dir_line = f"{empresa.get('direccion') or ''}, {empresa.get('comuna') or ''} - {empresa.get('ciudad') or ''}".strip(" ,-")
    if dir_line:
        par(dir_line.upper(), bold=True, size=8)
    if empresa.get("giro"):
        par(str(empresa["giro"]).upper(), bold=True, size=8)
    if empresa.get("telefono"):
        par(str(empresa["telefono"]), bold=True, size=8)
    par("", after=8)
    par("LIQUIDACION  DE  REMUNERACION", bold=True, size=11, align="center", subrayado=True, espaciado=60, after=6)
    par(C.mes_anio_es(periodo), bold=True, size=11, align="center", subrayado=True, after=4)
    par(str(trabajador.get("centro_costo") or liq.get("area") or "CASA MATRIZ").upper(), size=8, align="right", after=0)

    # ----- Ficha del trabajador (recuadro) -----
    nombre = f"{trabajador.get('apellido_paterno') or ''} {trabajador.get('apellido_materno') or ''} {trabajador.get('nombres') or ''}"
    nombre = " ".join(nombre.split()).upper()
    pactado = float(trabajador.get("pactado_salud_uf") or 0)
    es_isapre = (trabajador.get("salud") or "").upper() == "ISAPRE"
    pactado_txt = (f"{pactado:.4f}".rstrip("0").rstrip(".").replace(".", ",") + " UF") if (es_isapre and pactado) else "7,00  %"
    uf_txt = f"{uf:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    filas = [
        ("Código", str(trabajador.get("codigo") or trabajador.get("id") or ""), "Pactado Salud", pactado_txt),
        ("Nombre", nombre, "Base Tributable", _fmt_clp(liq.get("base_tributable"))),
        ("R.U.T.", trabajador.get("rut") or "", "U.F. del Mes", uf_txt),
        ("Fecha Ingreso", C.fecha_ddmmaaaa(liq.get("fecha_inicio")), "Tope Imponible", _fmt_clp(tope_afp)),
        ("Cargo", str(liq.get("cargo") or trabajador.get("cargo") or "").upper(), "Tope Imponible AFC", _fmt_clp(tope_afc)),
    ]
    info = doc.add_table(rows=len(filas), cols=6)
    _sin_bordes(info)
    for i, (a, b, c, d) in enumerate(filas):
        fila = info.rows[i].cells
        _txt(fila[0], a, size=9)
        _txt(fila[1], ":", size=9)
        _txt(fila[2], b, size=9)
        _txt(fila[3], c, size=9)
        _txt(fila[4], ":", size=9)
        _txt(fila[5], d, size=9, align="right")
        for j, cell in enumerate(fila):
            _cell_borders(cell, top=8 if i == 0 else None, left=6 if j == 0 else None,
                          right=6 if j == 5 else None)
    _anchos(info, [2.4, 0.5, 7.4, 3.4, 0.5, 4.4])

    # ----- Haberes / Descuentos -----
    detalle = liq.get("detalle") or []
    if isinstance(detalle, str):
        detalle = json.loads(detalle or "[]")
    dias = liq.get("dias_trabajados")
    dias_txt = f"{float(dias):g}" if dias is not None else "30"
    hab = [("SUELDO BASE", liq.get("sueldo_base") or liq.get("sueldo_calculado")),
           (f"SUELDO CALCULADO ({dias_txt})", liq.get("sueldo_calculado"))]
    imp_det = [d for d in detalle if d.get("tipo") == "Haber imponible"]
    if imp_det:
        hab += [(str(d["nombre"]).upper(), d["monto"]) for d in imp_det]
    elif float(liq.get("monto_horas_extras") or 0):
        hab.append(("HORAS EXTRAS", liq.get("monto_horas_extras")))
    if float(liq.get("gratificacion") or 0):
        hab.append(("GRATIFICACION LEGAL", liq.get("gratificacion")))
    if float(liq.get("otros_haberes") or 0):
        hab.append(("OTROS HABERES IMPONIBLES", liq.get("otros_haberes")))
    hab += [("", None), ("TOTAL IMPONIBLE", liq.get("total_imponible")), ("", None)]
    noimp = []
    for etiqueta, k in (("COLACION", "colacion"), ("MOVILIZACION", "movilizacion"),
                        ("ASIGNACION FAMILIAR", "asignacion_familiar")):
        if float(liq.get(k) or 0):
            noimp.append((etiqueta, liq.get(k)))
    noimp += [(str(d["nombre"]).upper(), d["monto"]) for d in detalle if d.get("tipo") == "Haber no imponible"]
    total_noimp = sum(float(m or 0) for _, m in noimp)
    if noimp:
        hab += noimp + [("", None), ("TOTAL NO IMPONIBLE", total_noimp)]

    afp_nombre = {"PlanVital": "PLAN VITAL"}.get(trabajador.get("afp"), (trabajador.get("afp") or "").upper())
    tasas = ind.get("afp_tasas") or {}
    if isinstance(tasas, str):
        tasas = json.loads(tasas or "{}")
    tasa = tasas.get(trabajador.get("afp"))
    salud = "FONASA" if not es_isapre else (trabajador.get("isapre") or "ISAPRE").upper()
    desc = []
    if float(liq.get("afp_monto") or 0) or trabajador.get("afp"):
        etiqueta = (f"{float(tasa):.2f}".replace(".", ",") + f"  % {afp_nombre}") if tasa else afp_nombre or "AFP"
        desc.append((etiqueta, liq.get("afp_monto")))
    desc.append((f" 7,00  % {salud}", liq.get("salud_monto")))
    if float(liq.get("adicional_isapre") or 0):
        desc.append(("ADICIONAL ISAPRE", liq.get("adicional_isapre")))
    if float(liq.get("afc_trabajador") or 0):
        desc.append(("SEGURO CESANTIA", liq.get("afc_trabajador")))
    total_impos = sum(float(liq.get(k) or 0) for k in ("afp_monto", "salud_monto", "adicional_isapre", "afc_trabajador"))
    desc += [("", None), ("TOTAL IMPOSICION", total_impos)]
    iu = float(liq.get("impuesto_unico") or 0)
    if iu:
        desc += [("", None), ("IMPUESTO UNICO", iu)]
    otros = []
    if float(liq.get("anticipo") or 0):
        otros.append(("ANTICIPO 1", liq.get("anticipo")))
    otros += [(str(d["nombre"]).upper(), d["monto"]) for d in detalle if d.get("tipo") == "Descuento"]
    if otros:
        desc += [("", None)] + otros + [("", None), ("TOTAL OTROS DESCUENTOS", sum(float(m or 0) for _, m in otros))]
    total_desc = float(liq.get("total_descuentos") or (total_impos + iu + sum(float(m or 0) for _, m in otros)))

    n = max(len(hab), len(desc)) + 1  # +1 fila en blanco al final del cuerpo
    hab += [("", None)] * (n - len(hab))
    desc += [("", None)] * (n - len(desc))
    filas_total = 1 + n + 1 + 1 + 1  # encabezado, cuerpo, totales, líquido, son
    tab = doc.add_table(rows=filas_total, cols=4)
    _sin_bordes(tab)
    LINEA = 6

    def bordes_fila(i, top=None, bottom=None):
        for j, cell in enumerate(tab.rows[i].cells):
            _cell_borders(cell, top=top, bottom=bottom, left=LINEA if j == 0 else None,
                          right=LINEA if j in (1, 3) else None)

    # encabezado
    _txt(tab.cell(0, 0), "HABERES", size=9, align="center", espaciado=60)
    _txt(tab.cell(0, 2), "DESCUENTOS", size=9, align="center", espaciado=60)
    tab.cell(0, 0).merge(tab.cell(0, 1))
    tab.cell(0, 2).merge(tab.cell(0, 3))
    bordes_fila(0, top=LINEA, bottom=LINEA)
    # cuerpo
    for k in range(n):
        i = 1 + k
        (h1, h2), (d1, d2) = hab[k], desc[k]
        _txt(tab.cell(i, 0), h1, size=9)
        _txt(tab.cell(i, 1), _fmt_clp(h2) if h2 is not None else "", size=9, align="right")
        _txt(tab.cell(i, 2), d1, size=9)
        _txt(tab.cell(i, 3), _fmt_clp(d2) if d2 is not None else "", size=9, align="right")
        bordes_fila(i)
    # totales
    it = 1 + n
    _txt(tab.cell(it, 0), "TOTAL HABERES      $", size=9, espaciado=50)
    _txt(tab.cell(it, 1), _fmt_clp(liq.get("total_haberes")), size=9, align="right")
    _txt(tab.cell(it, 2), "TOTAL DESCUENTOS     $", size=9, espaciado=50)
    _txt(tab.cell(it, 3), _fmt_clp(total_desc), size=9, align="right")
    bordes_fila(it, top=LINEA, bottom=LINEA)
    # líquido
    il = it + 1
    _txt(tab.cell(il, 0), "", size=9)
    _txt(tab.cell(il, 2), "LIQUIDO                    $", bold=True, size=11, espaciado=60)
    _txt(tab.cell(il, 3), _fmt_clp(liq.get("liquido")), bold=True, size=11, align="right")
    for j, cell in enumerate(tab.rows[il].cells):
        _cell_borders(cell, top=LINEA, left=LINEA if j == 0 else None, right=LINEA if j == 3 else None)
    # son
    isn = il + 1
    palabras = numero_a_palabras(int(round(float(liq.get("liquido") or 0))))
    tab.cell(isn, 0).merge(tab.cell(isn, 3))
    _txt(tab.cell(isn, 0), f"     SON:    {palabras} PESOS", size=9)
    _cell_borders(tab.cell(isn, 0), left=LINEA, right=LINEA, bottom=LINEA)
    for r in (il, isn):
        tab.rows[r].height = Cm(1.0)
    for r in (0, it):
        tab.rows[r].height = Cm(0.6)
    _anchos(tab, [7.2, 2.2, 7.0, 2.2])

    # ----- Certificación y firma -----
    par("", after=6)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(f"CERTIFICO QUE HE RECIBIDO DE  {str(empresa.get('razon_social') or '').upper()}\n"
                  "A MI ENTERA SATISFACCION, LA CANTIDAD INDICADA ANTERIORMENTE, COMO SALDO LIQUIDO DE MI SUELDO Y  "
                  "NO TENGO CARGO NI COBRO ALGUNO POSTERIOR QUE HACER, POR NINGUNO DE LOS MOTIVOS COMPRENDIDOS EN "
                  "ESTA LIQUIDACION.")
    r.font.name, r.font.size = FONT, Pt(8)
    for _ in range(4):
        par("", after=6)
    firma = doc.add_table(rows=3, cols=2)
    _sin_bordes(firma)
    _txt(firma.cell(0, 1), "_____________________________", size=9, align="center")
    _txt(firma.cell(1, 1), "FIRMA DEL TRABAJADOR", size=8, align="center")
    _txt(firma.cell(2, 1), "RECIBI COPIA", size=8, align="center")
    _anchos(firma, [11.0, 7.6])

    doc.save(ruta)
    return ruta


def generar_comprobante_feriado_docx(empresa, trabajador, vac, contrato, ruta):
    """
    Comprobante de Feriado (formato profesional chileno).
    vac: dict con fecha_inicio, fecha_termino, dias_habiles, dias_corridos, tipo, valor (opcional)
    contrato: dict con fecha_inicio y sueldo_base si existe
    """
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Cm(1.2)
        sec.bottom_margin = Cm(1.2)
        sec.left_margin = Cm(1.5)
        sec.right_margin = Cm(1.5)

    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(9)

    def _p(text, bold=False, size=9, center=False, right=False):
        p = doc.add_paragraph()
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif right:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.space_before = Pt(0)
        run = p.add_run(str(text))
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.bold = bold
        return p

    # Encabezado
    _p(str(empresa.get("razon_social") or "").upper(), bold=True, size=11, center=True)
    _p(str(empresa.get("rut") or ""), size=9, center=True)
    dir_line = f"{empresa.get('direccion') or ''}, {empresa.get('comuna') or ''} - {empresa.get('ciudad') or ''}".strip(" ,-")
    _p(dir_line.upper(), size=8, center=True)
    _p("ADMINISTRACION", size=9, right=True)

    # Título caja
    t = doc.add_table(rows=2, cols=3)
    _set_table_borders(t, sz="8")
    try:
        t.cell(0, 0).merge(t.cell(1, 0))
    except Exception:
        pass
    _cell(t.cell(0, 0), "COMPROBANTE\nDE FERIADO", bold=True, size=12, align="center")
    _cell(t.cell(0, 1), "LUGAR", bold=True, size=8, align="center")
    _cell(t.cell(0, 2), "FECHA EMISION", bold=True, size=8, align="center")
    lugar = (empresa.get("ciudad") or empresa.get("comuna") or "SANTIAGO").upper()
    fecha_em = C.fecha_larga()
    _cell(t.cell(1, 1), lugar, size=9, align="center")
    _cell(t.cell(1, 2), fecha_em, size=9, align="center")

    doc.add_paragraph()
    texto = (
        "En cumplimiento a las disposiciones legales vigentes se deja constancia que a contar de las fechas que se indican, el "
        "trabajador hara uso (Total o Parcial) del Feriado Anual con Remuneración integra de acuerdo al siguiente detalle:"
    )
    _p(texto, size=8)

    # Trabajador / RUT / Periodo
    nombre = f"{trabajador.get('nombres') or ''} {trabajador.get('apellido_paterno') or ''} {trabajador.get('apellido_materno') or ''}".strip().upper()
    info = doc.add_table(rows=2, cols=3)
    _set_table_borders(info, sz="4")
    _cell(info.cell(0, 0), "TRABAJADOR", bold=True, size=8)
    _cell(info.cell(0, 1), "RUT", bold=True, size=8, align="center")
    _cell(info.cell(0, 2), "PERIODO", bold=True, size=8, align="center")
    _cell(info.cell(1, 0), nombre, size=10)
    _cell(info.cell(1, 1), trabajador.get("rut") or "", size=10, align="center")
    _cell(info.cell(1, 2), "TOTAL", size=9, align="center")

    _p("DESCANSO EFECTIVO ENTRE LAS FECHAS QUE SE INDICAN:", bold=True, size=8)

    def _fmt_f(f):
        if not f:
            return ""
        s = str(f)[:10]
        if len(s) == 10 and s[4] == "-":
            return f"{s[8:10]}-{s[5:7]}-{s[0:4]}"
        return s

    f_ini = vac.get("fecha_inicio")
    f_fin = vac.get("fecha_termino")
    # Periodo contractual del contrato (año de feriado)
    c_ini = ""
    c_fin = ""
    if contrato:
        c_ini = contrato.get("fecha_inicio") or ""
        # Periodo feriado anual típico: año calendario del inicio de vacaciones
        try:
            y = int(str(f_ini)[:4]) if f_ini else date.today().year
            c_ini = f"01-01-{y}"
            c_fin = f"31-12-{y}"
        except Exception:
            pass

    fechas = doc.add_table(rows=2, cols=5)
    _set_table_borders(fechas, sz="4")
    _cell(fechas.cell(0, 0), "PERIODO CONTRACTUAL", bold=True, size=8)
    _cell(fechas.cell(0, 1), "DESDE:", size=8)
    _cell(fechas.cell(0, 2), _fmt_f(f_ini), size=9)
    _cell(fechas.cell(0, 3), "HASTA:", size=8)
    _cell(fechas.cell(0, 4), _fmt_f(f_fin), size=9)
    _cell(fechas.cell(1, 0), "", size=8)
    _cell(fechas.cell(1, 1), "DESDE:", size=8)
    _cell(fechas.cell(1, 2), c_ini if isinstance(c_ini, str) and "-" in str(c_ini)[-5:] else _fmt_f(c_ini), size=9)
    _cell(fechas.cell(1, 3), "HASTA:", size=8)
    _cell(fechas.cell(1, 4), c_fin, size=9)
    # Corregir etiquetas fila 0 = descanso efectivo
    _cell(fechas.cell(0, 0), "DESCANSO EFECTIVO", bold=True, size=8)
    _cell(fechas.cell(1, 0), "PERIODO CONTRACTUAL", bold=True, size=8)

    doc.add_paragraph()

    dias_hab = float(vac.get("dias_habiles") or 0)
    dias_prog = float(vac.get("dias_progresivas") or 0) if vac.get("tipo") == "Progresivas" else 0.0
    dias_adic = float(vac.get("dias_adicionales") or 0)
    domingo = float(vac.get("domingo_inhabiles") or 0)
    fraccionado = float(vac.get("feriado_fraccionado") or 0)
    saldo = float(vac.get("saldo_pendiente") or 0)
    total_dias = dias_hab + dias_prog + dias_adic

    # Valor: sueldo/30 * dias hábiles (práctica habitual)
    sueldo = 0.0
    if contrato:
        sueldo = float(contrato.get("sueldo_base") or 0)
    valor = float(vac.get("valor") or 0)
    if not valor and sueldo and total_dias:
        valor = round(sueldo / 30.0 * total_dias)

    # Detalle + totales lado a lado
    outer = doc.add_table(rows=1, cols=2)
    left = outer.cell(0, 0)
    right = outer.cell(0, 1)
    left.text = ""
    right.text = ""

    # Tabla detalle en celda izquierda
    det = left.add_table(rows=7, cols=2)
    _set_table_borders(det, sz="4")
    _cell(det.cell(0, 0), "DETALLE DE FERIADO", bold=True, size=8, align="center")
    _cell(det.cell(0, 1), "DIAS", bold=True, size=8, align="center")
    filas_d = [
        ("DIAS HABILES", f"{dias_hab:.2f}"),
        ("VACACIONES PROGRESIVAS", f"{dias_prog:.2f}"),
        ("DIAS ADICIONALES", f"{dias_adic:.2f}"),
        ("DOMINGO E INHABILES", f"{domingo:.2f}"),
        ("FERIADO FRACCIONADO", f"{fraccionado:.2f}" if fraccionado else ""),
        ("SALDO PENDIENTE", f"{saldo:.2f}"),
    ]
    for i, (lab, val) in enumerate(filas_d, 1):
        _cell(det.cell(i, 0), lab, size=8)
        _cell(det.cell(i, 1), val, size=9, align="right")

    # Tabla valor a la derecha
    val_t = right.add_table(rows=3, cols=2)
    _set_table_borders(val_t, sz="4")
    _cell(val_t.cell(0, 0), "DIAS", bold=True, size=8, align="center")
    _cell(val_t.cell(0, 1), "VALOR", bold=True, size=8, align="center")
    _cell(val_t.cell(1, 0), f"{total_dias:.2f}", size=11, align="center")
    _cell(val_t.cell(1, 1), _fmt_clp(valor), size=11, align="center")
    _cell(val_t.cell(2, 0), "TOTAL", bold=True, size=9, align="center")
    _cell(val_t.cell(2, 1), _fmt_clp(valor), bold=True, size=9, align="center")

    doc.add_paragraph()
    doc.add_paragraph()

    firmas = doc.add_table(rows=3, cols=2)
    _set_table_borders(firmas, sz="4")
    _cell(firmas.cell(0, 0), "", size=9)
    _cell(firmas.cell(0, 1), "", size=9)
    _cell(firmas.cell(1, 0), "", size=9)
    _cell(firmas.cell(1, 1), "", size=9)
    _cell(firmas.cell(2, 0), "NOMBRE Y FIRMA DEL EMPLEADOR O EMPRESA", bold=True, size=7, align="center")
    _cell(firmas.cell(2, 1), "NOMBRE Y FIRMA DEL TRABAJADOR", bold=True, size=7, align="center")

    doc.add_paragraph()
    nota = doc.add_paragraph()
    nr = nota.add_run(
        'NOTA:  Se deja constancia que el cálculo del feriado se ha hecho de conformidad a lo dispuesto en el capítulo VII, '
        '"Del Feriado anual y de los permisos, Capítulo 1 del Código del Trabajo".'
    )
    nr.font.name = "Arial"
    nr.font.size = Pt(7)

    doc.save(ruta)
    return ruta


def _formato_documento(doc):
    """Carta, márgenes 2,5 cm, Arial 11, texto justificado, interlineado 1,15."""
    for sec in doc.sections:
        sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
        sec.top_margin = sec.bottom_margin = Cm(2.2)
        sec.left_margin = sec.right_margin = Cm(2.5)
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    pf = st.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing = 1.15
    pf.space_after = Pt(8)


def _parrafo(doc, texto="", bold=False, size=11, align="justify", space_after=8, sangria_cm=0):
    p = doc.add_paragraph()
    p.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER, "left": WD_ALIGN_PARAGRAPH.LEFT,
                   "right": WD_ALIGN_PARAGRAPH.RIGHT}.get(align, WD_ALIGN_PARAGRAPH.JUSTIFY)
    if sangria_cm:
        p.paragraph_format.left_indent = Cm(sangria_cm)
    p.paragraph_format.space_after = Pt(space_after)
    if texto:
        r = p.add_run(texto)
        r.bold = bold
        r.font.name = "Arial"
        r.font.size = Pt(size)
    return p


def _clausula(doc, titulo, texto):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(f"{titulo} ")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r2 = p.add_run(texto)
    r2.font.name = "Arial"
    r2.font.size = Pt(11)
    return p


def _nombre(t):
    return f"{t.get('nombres') or ''} {t.get('apellido_paterno') or ''} {t.get('apellido_materno') or ''}".strip()


def _firmas(doc, empresa, trabajador, etiqueta_trab="TRABAJADOR"):
    tabla = doc.add_table(rows=2, cols=2)
    _sin_bordes(tabla)
    textos = [("_______________________________", "_______________________________"),
              (f"{empresa.get('razon_social')}\nRUT {empresa.get('rut')}\nEMPLEADOR",
               f"{_nombre(trabajador).upper()}\nRUN {trabajador.get('rut')}\n{etiqueta_trab}")]
    for i, fila in enumerate(textos):
        for j, t in enumerate(fila):
            cell = tabla.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(t)
            r.font.name, r.font.size = "Arial", Pt(10)
    _anchos(tabla, [8.3, 8.3])


def generar_contrato_docx(empresa, trabajador, contrato, ruta):
    doc = Document()
    _formato_documento(doc)
    _parrafo(doc, "CONTRATO DE TRABAJO", bold=True, size=14, align="center", space_after=12)

    p = _parrafo(doc, align="justify")
    p.add_run(f"En {empresa.get('comuna') or 'Santiago'}, a {C.fecha_larga()}, entre ")
    p.add_run(f"{empresa.get('razon_social')}, RUT {empresa.get('rut')}").bold = True
    p.add_run(f", representada por don/ña {empresa.get('representante_legal') or '__________'}, RUN "
              f"{empresa.get('rut_representante') or '__________'}, con domicilio en {empresa.get('direccion') or ''}, "
              f"comuna de {empresa.get('comuna') or ''}, en adelante «el empleador», y don/ña ")
    p.add_run(f"{_nombre(trabajador)}, RUN {trabajador.get('rut')}").bold = True
    p.add_run(f", de nacionalidad {trabajador.get('nacionalidad') or 'chilena'}, nacido/a el "
              f"{C.fecha_larga(trabajador.get('fecha_nacimiento')) if trabajador.get('fecha_nacimiento') else '__________'}, "
              f"domiciliado/a en {trabajador.get('direccion') or '__________'}, comuna de {trabajador.get('comuna') or '__________'}, "
              "en adelante «el trabajador», se ha convenido el siguiente contrato de trabajo:")

    _clausula(doc, "PRIMERO.-", f"El trabajador se obliga a desempeñar el cargo de {contrato.get('cargo')}, en "
              f"{contrato.get('lugar_trabajo') or empresa.get('direccion') or 'las dependencias del empleador'}, pudiendo ser "
              "trasladado a otro lugar o labor similar dentro de la misma ciudad, sin menoscabo para el trabajador (art. 12).")
    jornada = int(contrato.get("jornada_semanal") or C.JORNADA_DEFAULT)
    _clausula(doc, "SEGUNDO.-", f"La jornada ordinaria de trabajo será de {jornada} horas semanales, distribuidas "
              f"{contrato.get('horario') or 'de lunes a viernes'}. Las horas extraordinarias solo podrán trabajarse "
              "previo pacto escrito y se pagarán con un recargo del 50% sobre el sueldo convenido (art. 32).")

    lineas = [f"a) Sueldo base mensual: $ {C.fmt_clp(contrato.get('sueldo_base'))}."]
    tg = contrato.get("tipo_gratificacion") or "Monto fijo pactado"
    if tg.startswith("Art. 50"):
        lineas.append("b) Gratificación legal conforme al artículo 50 del Código del Trabajo: 25% de lo devengado en el "
                      "respectivo ejercicio por concepto de remuneraciones mensuales, con tope de 4,75 ingresos mínimos "
                      "mensuales, pagadera mensualmente como anticipo.")
    elif tg.startswith("Monto") and float(contrato.get("gratificacion") or 0) > 0:
        lineas.append(f"b) Gratificación mensual: $ {C.fmt_clp(contrato.get('gratificacion'))}.")
    if float(contrato.get("movilizacion") or 0) > 0:
        lineas.append(f"c) Asignación de movilización: $ {C.fmt_clp(contrato.get('movilizacion'))} (no imponible, art. 41).")
    if float(contrato.get("colacion") or 0) > 0:
        lineas.append(f"d) Asignación de colación: $ {C.fmt_clp(contrato.get('colacion'))} (no imponible, art. 41).")
    _clausula(doc, "TERCERO.-", "El empleador pagará al trabajador la siguiente remuneración, por períodos mensuales "
              "vencidos, el último día hábil de cada mes:")
    letras = "abcdefgh"
    for i, ln in enumerate(lineas):
        _parrafo(doc, letras[i] + ")" + ln[2:], sangria_cm=1.0, space_after=4)
    _parrafo(doc, "De las remuneraciones se deducirán los impuestos y cotizaciones de seguridad social que correspondan.")

    tipo = contrato.get("tipo_contrato") or "Indefinido"
    if tipo == "Plazo Fijo" and contrato.get("fecha_termino"):
        dur = f"El presente contrato es a plazo fijo y durará hasta el {C.fecha_larga(contrato.get('fecha_termino'))}."
    elif tipo == "Obra o faena":
        dur = "El presente contrato es por obra o faena y durará hasta la conclusión de la obra o servicio que le dio origen."
    else:
        dur = "El presente contrato es de duración indefinida."
    _clausula(doc, "CUARTO.-", dur)
    _clausula(doc, "QUINTO.-", f"El trabajador ingresó al servicio del empleador el {C.fecha_larga(contrato.get('fecha_inicio'))}. "
              f"Para efectos previsionales declara estar afiliado a AFP {trabajador.get('afp') or '______'} y al sistema de salud "
              f"{trabajador.get('salud') or '______'}{' (' + trabajador['isapre'] + ')' if trabajador.get('isapre') else ''}.")
    _clausula(doc, "SEXTO.-", "Se entienden incorporadas al presente contrato las disposiciones legales que se dicten y "
              "que regulen las relaciones laborales, así como el reglamento interno de orden, higiene y seguridad.")
    _parrafo(doc, "Para constancia, las partes firman el presente contrato en tres ejemplares, quedando uno en poder del trabajador.",
             align="justify", space_after=36)
    _firmas(doc, empresa, trabajador)
    doc.save(ruta)
    return ruta


def generar_finiquito_docx(empresa, trabajador, fin, contrato, ruta):
    """fin: resultado de finiquitos.calcular_finiquito (o fila de la tabla finiquitos)."""
    doc = Document()
    _formato_documento(doc)
    _parrafo(doc, "FINIQUITO DE CONTRATO DE TRABAJO", bold=True, size=14, align="center", space_after=12)
    nombre = _nombre(trabajador)

    p = _parrafo(doc, align="justify")
    p.add_run(f"En {empresa.get('comuna') or 'Santiago'}, a {C.fecha_larga()}, entre ")
    p.add_run(f"{empresa.get('razon_social')}, RUT {empresa.get('rut')}").bold = True
    p.add_run(f", representada legalmente por don/ña {empresa.get('representante_legal') or '__________'}, RUN "
              f"{empresa.get('rut_representante') or '__________'}, ambos domiciliados en {empresa.get('direccion') or ''}, "
              f"comuna de {empresa.get('comuna') or ''}, en adelante «el empleador», y don/ña ")
    p.add_run(f"{nombre}, RUN {trabajador.get('rut')}").bold = True
    p.add_run(f", de nacionalidad {trabajador.get('nacionalidad') or 'chilena'}, en adelante «el trabajador», "
              "se deja testimonio del siguiente finiquito:")

    _clausula(doc, "PRIMERO.-", f"El trabajador prestó servicios al empleador como {contrato.get('cargo')}, desde el "
              f"{C.fecha_larga(contrato.get('fecha_inicio'))} hasta el {C.fecha_larga(fin.get('fecha_termino'))}, fecha en "
              f"que su contrato terminó por la causal {fin.get('causal')} del Código del Trabajo.")
    _clausula(doc, "SEGUNDO.-", "El trabajador declara recibir en este acto, a su entera satisfacción, las siguientes sumas:")

    items = []
    if fin.get("indemnizacion_anos"):
        items.append((f"Indemnización por años de servicio ({fin.get('anos_pagar')} año(s) × $ {C.fmt_clp(fin.get('base_con_tope'))})",
                      fin["indemnizacion_anos"]))
    if fin.get("aviso_previo"):
        items.append(("Indemnización sustitutiva del aviso previo", fin["aviso_previo"]))
    if fin.get("vacaciones_proporcionales_monto"):
        dh = f"{float(fin.get('vacaciones_proporcionales_dias') or 0):.2f}".replace(".", ",")
        dc = f"{float(fin.get('feriado_dias_corridos') or 0):.2f}".replace(".", ",")
        items.append((f"Feriado legal y proporcional ({dh} días hábiles = {dc} días corridos)",
                      fin["vacaciones_proporcionales_monto"]))
    if fin.get("otros_montos"):
        items.append(("Otros haberes", fin["otros_montos"]))
    if fin.get("descuento_afc"):
        items.append(("(−) Aporte del empleador al Seguro de Cesantía (art. 13 Ley 19.728)", -fin["descuento_afc"]))
    if fin.get("otros_descuentos"):
        items.append(("(−) Otros descuentos", -fin["otros_descuentos"]))
    t = doc.add_table(rows=len(items) + 1, cols=2)
    _set_table_borders(t)
    for i, (lab, monto) in enumerate(items):
        _cell(t.cell(i, 0), lab, size=10)
        _cell(t.cell(i, 1), ("-$ " if monto < 0 else "$ ") + C.fmt_clp(abs(monto)), size=10, align="right")
    _cell(t.cell(len(items), 0), "TOTAL FINIQUITO", bold=True, size=10)
    _cell(t.cell(len(items), 1), "$ " + C.fmt_clp(fin.get("total_finiquito")), bold=True, size=10, align="right")
    t.autofit = False
    for i_col, ancho in enumerate((Cm(12.5), Cm(3.5))):
        t.columns[i_col].width = ancho
        for row_ in t.rows:
            row_.cells[i_col].width = ancho
    _parrafo(doc, "")

    total = int(round(float(fin.get("total_finiquito") or 0)))
    _clausula(doc, "TERCERO.-", f"El empleador paga en este acto al trabajador la suma de $ {C.fmt_clp(total)} "
              f"({numero_a_palabras(total)} PESOS), mediante transferencia electrónica a la cuenta "
              f"{trabajador.get('cuenta_banco') or '__________'} del {trabajador.get('banco') or '__________'}.")
    _clausula(doc, "CUARTO.-", "El empleador declara, y acredita con el certificado respectivo, que las cotizaciones "
              "previsionales y de salud del trabajador se encuentran íntegramente pagadas hasta el último día del mes "
              "anterior al del término de los servicios (art. 162 inciso quinto del Código del Trabajo).")
    _clausula(doc, "QUINTO.-", "El trabajador declara que durante la relación laboral recibió oportunamente el total de "
              "las remuneraciones, beneficios y demás prestaciones convenidas, y que, con el pago de las sumas "
              "indicadas, nada se le adeuda por concepto alguno, otorgando al empleador el más amplio, completo y "
              "total finiquito, salvo las reservas de derechos que se consignen a continuación.")
    _parrafo(doc, "Reserva de derechos del trabajador: _____________________________________________", size=10)
    _clausula(doc, "SEXTO.-", "Conforme al artículo 177 del Código del Trabajo, el presente finiquito solo podrá ser "
              "invocado por el empleador una vez firmado por el interesado y ratificado ante un ministro de fe "
              "(inspector del trabajo, notario, oficial del registro civil o secretario municipal), o suscrito "
              "electrónicamente en el portal de la Dirección del Trabajo.")
    _parrafo(doc, "Para constancia firman las partes en tres ejemplares.", space_after=36)
    _firmas(doc, empresa, trabajador)
    _parrafo(doc, "")
    _parrafo(doc, "RATIFICACIÓN ANTE MINISTRO DE FE: _______________________________   Fecha: ____________", size=9)
    doc.save(ruta)
    return ruta


def generar_anexo_docx(empresa, trabajador, contrato, anexo, ruta):
    """
    anexo: {"fecha", "vigencia", "cambios": {...}} con cambios posibles:
      sueldo_base, cargo, jornada_semanal (+horario), colacion, movilizacion, lugar_trabajo,
      duracion ("indefinido" o una fecha de prórroga), texto_libre.
    """
    cb = anexo.get("cambios") or {}
    doc = Document()
    _formato_documento(doc)
    _parrafo(doc, "ANEXO DE CONTRATO DE TRABAJO", bold=True, size=14, align="center", space_after=12)

    p = _parrafo(doc, align="justify")
    p.add_run(f"En {empresa.get('comuna') or 'Santiago'}, a {C.fecha_larga(anexo.get('fecha'))}, entre ")
    p.add_run(f"{empresa.get('razon_social')}, RUT {empresa.get('rut')}").bold = True
    p.add_run(f", representada por don/ña {empresa.get('representante_legal') or '__________'}, RUN "
              f"{empresa.get('rut_representante') or '__________'}, con domicilio en {empresa.get('direccion') or ''}, "
              f"comuna de {empresa.get('comuna') or ''}, en adelante «el empleador», y don/ña ")
    p.add_run(f"{_nombre(trabajador)}, RUN {trabajador.get('rut')}").bold = True
    p.add_run(f", en adelante «el trabajador», se conviene el siguiente anexo al contrato de trabajo de fecha "
              f"{C.fecha_larga(contrato.get('fecha_inicio'))}:")

    _clausula(doc, "PRIMERO.-", "Las partes acuerdan modificar el contrato de trabajo individualizado, a contar del "
              f"{C.fecha_larga(anexo.get('vigencia') or anexo.get('fecha'))}, en los siguientes términos:")
    lineas = []
    if cb.get("sueldo_base") is not None:
        lineas.append(f"Remuneración: el sueldo base mensual será de $ {C.fmt_clp(cb['sueldo_base'])} "
                      f"({numero_a_palabras(cb['sueldo_base']).lower()} pesos).")
    if cb.get("cargo"):
        lineas.append(f"Cargo: el trabajador desempeñará el cargo de {cb['cargo']}.")
    if cb.get("lugar_trabajo"):
        lineas.append(f"Lugar de trabajo: {cb['lugar_trabajo']}.")
    if cb.get("jornada_semanal"):
        lineas.append(f"Jornada: la jornada ordinaria será de {int(cb['jornada_semanal'])} horas semanales, distribuidas "
                      f"{cb.get('horario') or contrato.get('horario') or 'de lunes a viernes'}.")
    if cb.get("colacion") is not None:
        lineas.append(f"Asignación de colación: $ {C.fmt_clp(cb['colacion'])} mensuales (no imponible, art. 41).")
    if cb.get("movilizacion") is not None:
        lineas.append(f"Asignación de movilización: $ {C.fmt_clp(cb['movilizacion'])} mensuales (no imponible, art. 41).")
    if cb.get("duracion") == "indefinido":
        lineas.append("Duración: a contar de esta fecha el contrato de trabajo pasa a ser de duración indefinida.")
    elif cb.get("duracion"):
        lineas.append(f"Duración: el contrato a plazo fijo se prorroga hasta el {C.fecha_larga(cb['duracion'])}.")
    if (cb.get("texto_libre") or "").strip():
        lineas.append(cb["texto_libre"].strip())
    letras = "abcdefghij"
    for i, ln in enumerate(lineas):
        _parrafo(doc, f"{letras[i]}) {ln}", sangria_cm=1.0, space_after=6)

    _clausula(doc, "SEGUNDO.-", "En todo lo no modificado por el presente anexo, se mantienen plenamente vigentes las "
              "estipulaciones del contrato de trabajo y de sus anexos anteriores.")
    _clausula(doc, "TERCERO.-", "El presente anexo se firma en dos ejemplares del mismo tenor, quedando uno en poder "
              "de cada parte, y forma parte integrante del contrato de trabajo.")
    _parrafo(doc, "", space_after=36)
    _firmas(doc, empresa, trabajador)
    doc.save(ruta)
    return ruta

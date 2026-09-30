import csv

import pytest

from remu import db, libros, previred, seguridad
from remu import procesos as PR


@pytest.fixture
def empresa_mov(base_temporal):
    conn = db.get_conn()
    eid = db.insert(conn, "empresas", dict(rut="76.065.376-4", razon_social="EMPRESA RRHH SPA", mutual="ACHS",
                                           tasa_mutual=0.93, region_codigo=13, comuna_codigo=13101))
    tid = db.insert(conn, "trabajadores", dict(empresa_id=eid, rut="12.345.678-5", nombres="JUAN", apellido_paterno="PEREZ",
                                               apellido_materno="SOTO", afp="Habitat", salud="FONASA",
                                               fecha_nacimiento="1990-01-01"))
    db.insert(conn, "contratos", dict(trabajador_id=tid, empresa_id=eid, cargo="VENDEDOR", fecha_inicio="2024-01-01",
                                      tipo_contrato="Indefinido", sueldo_base=840_000, colacion=60_000, movilizacion=30_000,
                                      tipo_gratificacion="Art. 50 (25% con tope 4,75 IMM)", jornada_semanal=42))
    cb = db.insert(conn, "conceptos", dict(empresa_id=eid, nombre="Bono metas", tipo="Haber imponible", codigo_lre=2113))
    cv = db.insert(conn, "conceptos", dict(empresa_id=eid, nombre="Viático", tipo="Haber no imponible", codigo_lre=2303))
    cd = db.insert(conn, "conceptos", dict(empresa_id=eid, nombre="Préstamo empresa", tipo="Descuento", codigo_lre=3183))
    conn.commit()
    emp = db.row(conn, "SELECT * FROM empresas WHERE id=?", (eid,))
    # mes anterior sin movimientos (para la renta imponible de la licencia)
    PR.calcular_periodo(conn, emp, "2026-07", db.get_indicadores("2026-07", conn))
    PR.guardar_movimiento(conn, eid, tid, "2026-08", dict(
        dias_trabajados=25, licencia=5, licencia_desde="2026-08-10", licencia_hasta="2026-08-14",
        anticipo=100_000, aguinaldo=50_000, bono_desempeno=40_000, cant_he_50=10, cant_he_100=2, cant_hd=8),
        {cb: 25_000, cv: 15_000, cd: 20_000}, "supervisor")
    conn.commit()
    PR.calcular_periodo(conn, emp, "2026-08", db.get_indicadores("2026-08", conn))
    liq = db.row(conn, "SELECT * FROM liquidaciones WHERE periodo='2026-08'")
    conn.close()
    return emp, liq


def test_movimiento_alimenta_liquidacion(empresa_mov):
    _, l = empresa_mov
    hora = 840_000 / (42 * 30 / 7)
    he = round(hora * 1.5 * 10) + round(hora * 2.0 * 2) + round(hora * 0.3 * 8)
    assert l["dias_trabajados"] == 25 and l["dias_licencia"] == 5
    assert l["sueldo_calculado"] == 700_000
    assert l["monto_horas_extras"] == he
    assert l["aguinaldo"] == 50_000 and l["bonos_imponibles"] == 40_000 + 25_000
    assert l["total_imponible"] == 700_000 + he + 50_000 + 65_000 + l["gratificacion"]
    assert l["haberes_no_imponibles"] == 15_000
    assert l["anticipo"] == 100_000 and l["otros_descuentos"] == 20_000
    assert l["colacion"] == 50_000 and l["movilizacion"] == 25_000  # contrato prorrateado a 25 días
    assert l["liquido"] == l["total_haberes"] - l["total_descuentos"]


def test_lre_usa_codigos_de_conceptos(empresa_mov, tmp_path):
    emp, _ = empresa_mov
    ruta = libros.generar_lre_csv(emp["id"], "2026-08", str(tmp_path / "lre.csv"))
    with open(ruta, encoding="utf-8") as f:
        filas = list(csv.reader(f, delimiter=";"))
    h = {x.split("(")[-1].rstrip(")"): i for i, x in enumerate(filas[0])}
    r = filas[1]
    assert int(r[h["2110"]]) == 50_000 and int(r[h["2113"]]) == 65_000
    assert int(r[h["2303"]]) == 15_000 and int(r[h["3183"]]) == 20_000 and int(r[h["3188"]]) == 100_000
    assert int(r[h["2107"]]) > 0 and int(r[h["1116"]]) == 5
    assert int(r[h["5501"]]) == int(r[h["5201"]]) - int(r[h["5301"]])


def test_previred_licencia(empresa_mov, tmp_path):
    emp, _ = empresa_mov
    res = previred.generar_previred_txt(emp["id"], "2026-08", str(tmp_path / "p.txt"))
    f = open(res["ruta"], encoding="latin-1").read().strip().split(";")
    assert f[14] == "3" and f[15] == "10-08-2026" and f[16] == "14-08-2026"
    assert int(f[91]) > 0  # renta imponible mes anterior proporcional


def test_asiento_cuadra_con_movimientos(empresa_mov, tmp_path):
    from openpyxl import load_workbook
    emp, _ = empresa_mov
    ruta = libros.generar_libro_remuneraciones(emp["id"], "2026-08", str(tmp_path / "l.xlsx"))
    tot = [r for r in load_workbook(ruta)["Asiento Centralizacion"].iter_rows(values_only=True) if r[0] == "TOTALES"][0]
    assert tot[1] == tot[2]


def test_estado_periodo(empresa_mov):
    emp, _ = empresa_mov
    conn = db.get_conn()
    assert PR.estado_periodo(conn, emp["id"], "2026-08")["estado"] == "Abierto"
    PR.guardar_estado_periodo(conn, emp["id"], "2026-08", "Enviado", "cliente")
    conn.commit()
    assert PR.estado_periodo(conn, emp["id"], "2026-08")["estado"] == "Enviado"
    conn.close()


def test_modulos_usuario(base_temporal):
    seguridad.crear_usuario("admin", "Clave2026x", rol="admin")
    seguridad.crear_usuario("cliente", "Demo2026x", empresas=[1], modulos=["rrhh"])
    u, _ = seguridad.autenticar("cliente", "Demo2026x")
    assert seguridad.modulos_usuario(u) == {"rrhh"}
    a, _ = seguridad.autenticar("admin", "Clave2026x")
    assert seguridad.modulos_usuario(a) == {"rrhh", "remuneraciones"}

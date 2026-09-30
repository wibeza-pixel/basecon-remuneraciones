import csv
from datetime import date

import pytest
from openpyxl import load_workbook

from remu import config as C
from remu import db, libros, previred, seguridad
from remu.procesos import calcular_periodo


@pytest.fixture
def empresa(base_temporal):
    conn = db.get_conn()
    eid = db.insert(conn, "empresas", dict(rut="76.065.376-4", razon_social="EMPRESA PRUEBA SPA", comuna="Santiago",
                                           region_codigo=13, comuna_codigo=13101, mutual="ACHS", tasa_mutual=0.93,
                                           caja_compensacion="LOS ANDES"))
    t1 = db.insert(conn, "trabajadores", dict(empresa_id=eid, rut="12.345.678-5", nombres="JUAN", apellido_paterno="PEREZ",
                                              apellido_materno="SOTO", afp="Habitat", salud="FONASA", numero_cargas=1,
                                              tramo_asignacion_familiar="B", fecha_nacimiento="1990-01-01"))
    t2 = db.insert(conn, "trabajadores", dict(empresa_id=eid, rut="11.111.111-1", nombres="ANA", apellido_paterno="ROJAS",
                                              apellido_materno="DIAZ", afp="Provida", salud="ISAPRE", isapre="COLMENA",
                                              pactado_salud_uf=8.0, sexo="F", fecha_nacimiento="1985-05-05"))
    db.insert(conn, "contratos", dict(trabajador_id=t1, empresa_id=eid, cargo="ANALISTA", fecha_inicio="2026-08-01",
                                      tipo_contrato="Indefinido", sueldo_base=900_000,
                                      tipo_gratificacion="Art. 50 (25% con tope 4,75 IMM)", colacion=40_000, jornada_semanal=42))
    db.insert(conn, "contratos", dict(trabajador_id=t2, empresa_id=eid, cargo="GERENTE", fecha_inicio="2020-01-01",
                                      tipo_contrato="Indefinido", sueldo_base=3_500_000,
                                      tipo_gratificacion="Monto fijo pactado", gratificacion=219_115, jornada_semanal=42))
    conn.commit()
    emp = db.row(conn, "SELECT * FROM empresas WHERE id=?", (eid,))
    ind = db.get_indicadores("2026-08", conn)
    n, _ = calcular_periodo(conn, emp, "2026-08", ind)
    conn.close()
    assert n == 2
    return emp


def test_init_no_pisa_indicadores(base_temporal):
    conn = db.get_conn()
    db.guardar_indicadores(conn, "2026-08", {"uf": 12345.0})
    conn.commit()
    conn.close()
    db._init_done = False
    db.init_db()
    assert db.get_indicadores("2026-08")["uf"] == 12345.0


def test_liquidaciones_guardan_impuesto(empresa):
    conn = db.get_conn()
    ls = db.rows(conn, "SELECT * FROM liquidaciones")
    conn.close()
    gerente = [x for x in ls if x["sueldo_base"] == 3_500_000][0]
    assert gerente["impuesto_unico"] > 0
    assert gerente["liquido"] == gerente["total_haberes"] - gerente["total_descuentos"]


def test_previred_105_campos(empresa, tmp_path):
    res = previred.generar_previred_txt(empresa["id"], "2026-08", str(tmp_path / "p.txt"))
    assert res["lineas"] == 2 and not res["errores"]
    lineas = open(res["ruta"], encoding="latin-1").read().strip().splitlines()
    for ln in lineas:
        f = ln.split(";")
        assert len(f) == 105
    juan = [ln.split(";") for ln in lineas if ln.startswith("12345678;")][0]
    assert juan[14] == "1"            # campo 15: contratación indefinida en el mes
    assert juan[15] == "01-08-2026"   # campo 16
    assert juan[25] == "05"           # AFP Habitat
    assert juan[74] == "07"           # Fonasa
    assert juan[82] == "01"           # CCAF Los Andes
    assert int(juan[89]) > 0          # parte CCAF del 7%
    assert juan[95] == "01"           # ACHS
    ana = [ln.split(";") for ln in lineas if ln.startswith("11111111;")][0]
    assert ana[74] == "04" and int(ana[80]) > 0  # Colmena + adicional


def test_lre_estructura(empresa, tmp_path):
    ruta = libros.generar_lre_csv(empresa["id"], "2026-08", str(tmp_path / "lre.csv"))
    with open(ruta, encoding="utf-8") as f:
        filas = list(csv.reader(f, delimiter=";"))
    assert filas[0][0] == "Rut trabajador(1101)"
    assert len(filas[0]) == len(libros.LRE_COLUMNAS) == len(filas[1])
    h = {x.split("(")[-1].rstrip(")"): i for i, x in enumerate(filas[0])}
    for fila in filas[1:]:
        assert int(fila[h["5501"]]) == int(fila[h["5201"]]) - int(fila[h["5301"]])
    assert any(int(f[h["3161"]]) > 0 for f in filas[1:])


def test_asiento_cuadra_con_impuesto(empresa, tmp_path):
    ruta = libros.generar_libro_remuneraciones(empresa["id"], "2026-08", str(tmp_path / "libro.xlsx"))
    ws = load_workbook(ruta)["Asiento Centralizacion"]
    textos = [str(c.value) for c in ws["A"] if c.value]
    assert any("IMPUESTO UNICO" in t for t in textos)
    tot = [r for r in ws.iter_rows(values_only=True) if r[0] == "TOTALES"][0]
    assert tot[1] == tot[2] and tot[3] == "Debe = Haber"


def test_1887_informa_impuesto(empresa, tmp_path):
    conn = db.get_conn()
    conn.execute("UPDATE liquidaciones SET periodo='2025-08'")
    conn.commit()
    conn.close()
    ruta = libros.generar_1887_csv_sii(empresa["id"], 2026, str(tmp_path / "dj.csv"))
    with open(ruta, encoding="utf-8") as f:
        filas = list(csv.DictReader(f, delimiter=";"))
    assert sum(int(x["IMPUESTO_UNICO_RETENIDO"]) for x in filas) > 0


def test_usuarios(base_temporal):
    assert not seguridad.hay_usuarios()
    with pytest.raises(ValueError):
        seguridad.crear_usuario("x", "123")
    seguridad.crear_usuario("admin", "Clave2026x", rol="admin")
    uid = seguridad.crear_usuario("cliente", "Demo2026x", empresas=[1], dias_acceso=15)
    assert seguridad.autenticar("admin", "mala")[0] is None
    u, _ = seguridad.autenticar("cliente", "Demo2026x")
    assert u and u["dias_restantes"] == 15 and seguridad.empresas_permitidas(u) == {1}
    seguridad.actualizar_usuario(uid, fecha_expira=date(2020, 1, 1))
    u, motivo = seguridad.autenticar("cliente", "Demo2026x")
    assert u is None and "finalizado" in motivo


def test_firma_detecta_edicion_manual(base_temporal, monkeypatch):
    monkeypatch.setenv("BASECON_SECRET", "secreto-de-prueba")
    uid = seguridad.crear_usuario("cliente", "Demo2026x", empresas=[1], dias_acceso=15)
    conn = db.get_conn()
    conn.execute("UPDATE usuarios SET dias_acceso=9999 WHERE id=?", (uid,))
    conn.commit()
    conn.close()
    u, motivo = seguridad.autenticar("cliente", "Demo2026x")
    assert u is None and "modificados" in motivo


def test_rut_valido():
    assert C.rut_valido("12.345.678-5") and not C.rut_valido("12.345.678-9")
